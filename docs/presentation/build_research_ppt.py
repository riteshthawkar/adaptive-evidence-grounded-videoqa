from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
REPORT_DIR = OUT_DIR.parent / "report"
PPTX_PATH = OUT_DIR / "research_presentation_8min.pptx"
PPTX_ALIAS_PATH = OUT_DIR / "research_presentation.pptx"
TRADEOFF_PATH = OUT_DIR / "qwen_test_tradeoff_presentation.png"
GROUNDING_PATH = OUT_DIR / "main_comparison_grounding_presentation.png"


BLUE = RGBColor(42, 89, 143)
BLUE_DARK = RGBColor(25, 55, 92)
BLUE_LIGHT = RGBColor(232, 240, 250)
GOLD = RGBColor(176, 132, 0)
GREEN = RGBColor(59, 139, 116)
GRAY = RGBColor(91, 99, 112)
LIGHT_GRAY = RGBColor(244, 246, 249)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 24, 31)


def make_figures() -> None:
    methods = ["One segment", "Learned two-item", "Fixed 3+3"]
    costs = [1.500, 2.795, 7.500]
    acc = [0.668, 0.696, 0.724]
    gqa = [0.221, 0.297, 0.454]
    iop = [0.311, 0.410, 0.615]
    colors = ["#2a598f", "#3d8b74", "#b08400"]

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.edgecolor": "#53606f",
            "axes.labelcolor": "#1a2433",
            "xtick.color": "#1a2433",
            "ytick.color": "#1a2433",
        }
    )

    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.2), dpi=220)
    panels = [
        (axes[0], acc, "Answer accuracy", "Acc@QA"),
        (axes[1], gqa, "Grounded answer accuracy", "Acc@GQA"),
    ]
    for ax, vals, title, ylabel in panels:
        ax.plot(costs, vals, color="#53606f", linewidth=1.8, zorder=1)
        ax.scatter(costs, vals, c=colors, s=115, edgecolor="white", linewidth=1.5, zorder=2)
        for x, y, label in zip(costs, vals, methods):
            ax.annotate(
                f"{label}\n{y:.3f}",
                (x, y),
                xytext=(7, 6),
                textcoords="offset points",
                fontsize=8.5,
                color="#1a2433",
            )
        ax.set_title(title, fontsize=12.5, fontweight="bold", pad=10)
        ax.set_xlabel("Evidence cost")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0.8, 8.15)
        ax.grid(True, axis="both", alpha=0.25)
    fig.suptitle("Current Qwen Test Summary: Accuracy-Cost-Grounding Tradeoff", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.92])
    fig.savefig(TRADEOFF_PATH, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    names = ["Temp[CLIP]", "FrozenBiLM", "SeViLA", "QGAC-TR", "One seg. (Our)", "Learned (Our)", "Fixed 3+3 (Our)"]
    acc_q = [0.602, 0.708, 0.681, 0.636, 0.668, 0.696, 0.724]
    acc_g = [0.160, 0.175, 0.166, 0.183, 0.221, 0.297, 0.454]
    x = range(len(names))
    fig, ax = plt.subplots(figsize=(11.5, 4.2), dpi=220)
    ax.bar([i - 0.18 for i in x], acc_q, width=0.36, label="Acc@QA", color="#7b9bc2")
    ax.bar([i + 0.18 for i in x], acc_g, width=0.36, label="Acc@GQA", color="#3d8b74")
    for i, value in enumerate(acc_g):
        ax.text(i + 0.18, value + 0.014, f"{value:.3f}", ha="center", va="bottom", fontsize=7.8)
    ax.axvline(3.5, color="#53606f", linestyle="--", linewidth=1.0)
    ax.text(1.7, 0.76, "Published methods", ha="center", color="#53606f", fontsize=9.5)
    ax.text(5.25, 0.76, "Our operating points", ha="center", color="#53606f", fontsize=9.5)
    ax.set_ylim(0, 0.82)
    ax.set_ylabel("Score")
    ax.set_title("Answer Accuracy and Grounded Accuracy Diverge", fontsize=14, fontweight="bold", pad=10)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=20, ha="right")
    ax.grid(True, axis="y", alpha=0.25)
    ax.legend(frameon=False, loc="upper left", ncol=2)
    fig.tight_layout()
    fig.savefig(GROUNDING_PATH, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_textbox(slide, x, y, w, h, text, font_size=20, color=BLACK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_header(slide, title, slide_no, total):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    top = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.52))
    top.fill.solid()
    top.fill.fore_color.rgb = BLUE
    top.line.fill.background()
    add_textbox(slide, 0.35, 0.075, 10.6, 0.36, title, font_size=18, color=WHITE, bold=True)
    add_textbox(slide, 11.55, 0.095, 1.35, 0.3, f"{slide_no} / {total}", font_size=18, color=WHITE, align=PP_ALIGN.RIGHT)
    bottom = slide.shapes.add_shape(1, Inches(0), Inches(7.19), Inches(13.333), Inches(0.31))
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = BLUE_LIGHT
    bottom.line.fill.background()
    add_textbox(slide, 0.35, 7.25, 8.5, 0.18, "Adaptive Evidence Acquisition for Grounded VideoQA", font_size=8.5, color=BLUE_DARK)
    add_textbox(slide, 10.25, 7.25, 2.7, 0.18, "Final Presentation", font_size=8.5, color=BLUE_DARK, align=PP_ALIGN.RIGHT)


def add_block(slide, x, y, w, h, title, body, accent=BLUE, body_size=17, title_size=15):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE_LIGHT
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.2)
    add_textbox(slide, x + 0.18, y + 0.12, w - 0.36, 0.28, title, font_size=title_size, color=accent, bold=True)
    add_textbox(slide, x + 0.18, y + 0.52, w - 0.36, h - 0.64, body, font_size=body_size, color=BLACK)
    return shape


def add_bullets(slide, x, y, w, h, items, font_size=19, color=BLACK, bullet=True, level=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = level
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
        if bullet:
            p._p.get_or_add_pPr().set("marL", "285750")
    return box


def set_cell(cell, text, font_size=10, bold=False, fill=None, color=BLACK, align=PP_ALIGN.CENTER):
    cell.text = text
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color


def add_table(slide, x, y, w, h, rows, headers, font_size=10, highlight_start=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for c, header in enumerate(headers):
        set_cell(table.cell(0, c), header, font_size=font_size, bold=True, fill=BLUE, color=WHITE)
    for r, row in enumerate(rows, start=1):
        fill = WHITE if r % 2 else LIGHT_GRAY
        if highlight_start is not None and r >= highlight_start:
            fill = RGBColor(235, 244, 241)
        for c, val in enumerate(row):
            align = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            set_cell(table.cell(r, c), str(val), font_size=font_size, fill=fill, align=align)
    return table_shape


def add_metric_cards(slide, cards, x=0.65, y=5.75, w=12.0, h=0.95):
    card_w = w / len(cards)
    for i, (label, value, color) in enumerate(cards):
        shape = slide.shapes.add_shape(1, Inches(x + i * card_w + 0.05), Inches(y), Inches(card_w - 0.1), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        shape.line.color.rgb = color
        shape.line.width = Pt(1.4)
        add_textbox(slide, x + i * card_w + 0.14, y + 0.14, card_w - 0.28, 0.28, label, font_size=10.5, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + i * card_w + 0.14, y + 0.44, card_w - 0.28, 0.36, value, font_size=18, color=color, bold=True, align=PP_ALIGN.CENTER)


def build_deck() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 15

    # 1
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.58))
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    add_textbox(slide, 0.72, 1.08, 11.9, 1.05, "Adaptive Evidence Acquisition for\nGrounded Video Question Answering", font_size=32, color=BLUE_DARK, bold=True)
    add_textbox(slide, 0.75, 2.55, 11.5, 0.52, "A cost-aware evidence selection view of grounded VideoQA", font_size=20, color=GRAY)
    add_block(slide, 0.75, 3.45, 5.75, 1.45, "Opening question", "How much visual evidence is enough to answer a video question while staying temporally grounded?", accent=GREEN, body_size=18)
    add_block(slide, 6.85, 3.45, 5.75, 1.45, "Main claim", "Answer accuracy, evidence cost, and temporal grounding must be evaluated together.", accent=GOLD, body_size=18)
    add_textbox(slide, 0.75, 6.18, 11.8, 0.35, "Aniket Salunke, Ketan More, Nazish Baliyan, Ritesh Thawkar | MBZUAI", font_size=13, color=GRAY)

    # 2
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Three Questions We Answer", 2, total)
    questions = [
        "Q1. Do all video questions need the same amount of visual evidence?",
        "Q2. Can a compact learned policy preserve answer accuracy at lower cost?",
        "Q3. Is poor performance caused by weak evidence selection or by a weak answerer?",
    ]
    y_positions = [1.05, 2.55, 4.05]
    colors = [BLUE, GREEN, GOLD]
    for i, (q, y, c) in enumerate(zip(questions, y_positions, colors), start=1):
        add_block(slide, 0.9, y, 11.55, 1.02, f"Research Question {i}", q, accent=c, body_size=20)
    add_textbox(slide, 1.05, 5.85, 11.2, 0.52, "We answer these after defining the task, method, and evidence-budget experiments.", font_size=18, color=GRAY, align=PP_ALIGN.CENTER)

    # 3
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Motivation: Correct Answers Can Be Ungrounded", 3, total)
    add_bullets(
        slide,
        0.8,
        1.0,
        5.8,
        3.0,
        [
            "VideoQA evidence is distributed over time.",
            "A correct option may be guessed from priors or broad scene context.",
            "Fixed frame/clip budgets can waste evidence or miss the support moment.",
        ],
        font_size=20,
    )
    add_block(slide, 7.0, 1.0, 5.35, 2.2, "Grounded VideoQA asks two things", "1. Is the answer correct?\n2. Does the selected evidence overlap the annotated support moment?", accent=BLUE, body_size=18)
    add_block(slide, 7.0, 3.7, 5.35, 1.65, "Why it matters", "Grounding turns VideoQA from answer classification into evidence-supported reasoning.", accent=GREEN, body_size=18)
    add_metric_cards(slide, [("Failure mode", "correct but weakly grounded", GOLD), ("Desired behavior", "correct and grounded", GREEN)], x=1.0, y=5.78, w=11.3, h=0.85)

    # 4
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Related Work and Gap", 4, total)
    add_table(
        slide,
        0.65,
        1.0,
        12.05,
        3.65,
        [
            ["VideoQA benchmarks", "MovieQA, TGIF-QA, TVQA, NExT-QA", "Temporal and causal reasoning"],
            ["Grounded VideoQA", "TVQA+, NExT-GQA", "Answer must be supported by temporal evidence"],
            ["Pretrained VLMs", "CLIP, FrozenBiLM, SeViLA, Qwen2.5-VL", "Strong visual-language reasoning"],
            ["Retrieval-augmented video", "VideoRAG-style selective context", "Long-video evidence construction"],
        ],
        ["Area", "Examples", "What they provide"],
        font_size=12,
    )
    add_block(slide, 0.9, 5.05, 11.55, 1.05, "Gap", "Prior work often fixes the amount of context or reports answer quality without jointly charging evidence cost and measuring temporal support.", accent=GOLD, body_size=18)

    # 5
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Task Formulation", 5, total)
    add_block(slide, 0.75, 0.95, 5.75, 1.45, "Grounded VideoQA example", "x = (v, q, A, y, T)\nvideo, question, answer options, label, support span", accent=BLUE, body_size=17)
    add_block(slide, 6.85, 0.95, 5.75, 1.45, "Evidence pool", "E(x) = {e1, ..., eN}\nframes and short temporal segments with cost and interval metadata", accent=GREEN, body_size=17)
    add_block(slide, 0.75, 2.85, 5.75, 1.45, "Policy", "Selects S subset E(x)\nC(S) = sum_i c_i", accent=GOLD, body_size=19)
    add_block(slide, 6.85, 2.85, 5.75, 1.45, "Answerer", "Predicts answer from question, options, and selected evidence only.", accent=BLUE, body_size=18)
    add_textbox(slide, 1.0, 5.45, 11.3, 0.48, "Optimization intuition: maximize answer accuracy and grounding while minimizing acquired evidence cost.", font_size=19, color=BLUE_DARK, bold=True, align=PP_ALIGN.CENTER)

    # 6
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Method: Adaptive Evidence Acquisition Pipeline", 6, total)
    pipeline = REPORT_DIR / "figure_1.png"
    if pipeline.exists():
        slide.shapes.add_picture(str(pipeline), Inches(1.28), Inches(0.72), width=Inches(10.8))
    else:
        add_bullets(slide, 1.0, 1.0, 11.0, 2.0, ["Normalize NExT-GQA", "Construct evidence pool", "Select evidence", "Answer", "Evaluate"], font_size=24)
    add_block(
        slide,
        0.85,
        5.58,
        5.75,
        1.02,
        "Candidate pool",
        "Sample frames and 4s sliding segments; rank top candidates with CLIP question-option similarity.",
        accent=BLUE,
        body_size=12,
        title_size=11,
    )
    add_block(
        slide,
        6.75,
        5.58,
        5.75,
        1.02,
        "Controlled answerer swap",
        "Keep selected evidence fixed, then swap only the answerer to separate evidence quality from reasoning quality.",
        accent=GREEN,
        body_size=12,
        title_size=11,
    )

    # 7
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Evidence Policies and Cost Ladder", 7, total)
    add_table(
        slide,
        0.7,
        0.95,
        11.95,
        3.0,
        [
            ["One segment", "highest-scoring segment", "1", "1.500", "Low-cost baseline"],
            ["Learned two-item", "linear router: frame/segment/stop", "2", "2.795", "Efficient middle point"],
            ["Fixed 3+3", "3 frames + 3 segments", "6", "7.500", "High-coverage control"],
            ["Fixed 6+6", "6 frames + 6 segments", "about 12", "14.908", "Large frozen-answerer control"],
        ],
        ["Policy", "Selection", "Count", "Cost", "Role"],
        font_size=11.5,
        highlight_start=2,
    )
    add_block(slide, 1.0, 4.55, 5.25, 1.35, "Policy training", "The learned linear router is trained from oracle traces; Qwen is not trained or used for routing.", accent=BLUE, body_size=16)
    add_block(slide, 7.0, 4.55, 5.25, 1.35, "Why fixed controls matter", "They show what grounding improves when we spend much more visual context.", accent=GREEN, body_size=16)

    # 8
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Experimental Setup and Metrics", 8, total)
    add_block(slide, 0.75, 0.95, 3.75, 1.55, "Dataset", "NExT-GQA\nValidation: 3,358 examples\nTest: 5,553 examples", accent=BLUE, body_size=16)
    add_block(slide, 4.8, 0.95, 3.75, 1.55, "Answerers", "Frozen CLIP-style scorer\nQwen2.5-VL-3B-Instruct", accent=GREEN, body_size=16)
    add_block(slide, 8.85, 0.95, 3.75, 1.55, "Protocol", "Same selected evidence\nSwap only final answerer", accent=GOLD, body_size=16)
    add_table(
        slide,
        0.85,
        3.05,
        11.65,
        2.45,
        [
            ["Acc@QA", "Whether the answer letter is correct"],
            ["Cost / Count", "Amount of visual evidence consumed"],
            ["mIoP, IoP@0.5", "How much selected evidence lies inside support span"],
            ["mIoU, IoU@0.5", "Stricter temporal overlap with support span"],
            ["Acc@GQA", "Correct answer plus IoP@0.5 grounding"],
        ],
        ["Metric", "Meaning"],
        font_size=12.5,
    )

    # 9
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Main Comparison: Answer Accuracy Is Not Grounding", 9, total)
    slide.shapes.add_picture(str(GROUNDING_PATH), Inches(0.75), Inches(0.85), width=Inches(11.85))
    add_block(
        slide,
        0.9,
        5.48,
        3.75,
        1.15,
        "Context",
        "Prior rows provide scale; our grounding uses selected-evidence best overlap.",
        accent=BLUE,
        body_size=13,
        title_size=12,
    )
    add_block(
        slide,
        4.8,
        5.48,
        3.75,
        1.15,
        "Our proposed point",
        "Learned evidence gives the best cost-aware tradeoff.",
        accent=GREEN,
        body_size=13,
        title_size=12,
    )
    add_block(
        slide,
        8.7,
        5.48,
        3.75,
        1.15,
        "Our upper point",
        "Fixed 3+3 gives the strongest grounding.",
        accent=GOLD,
        body_size=13,
        title_size=12,
    )

    # 10
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Result 1: Frozen Answerer Is the Bottleneck", 10, total)
    add_table(
        slide,
        0.75,
        1.0,
        11.85,
        2.65,
        [
            ["One segment", "0.378", "1.500", "0.311", "0.124"],
            ["Learned two-item", "0.381", "2.795", "0.410", "0.164"],
            ["Fixed 3+3", "0.388", "7.500", "0.615", "0.245"],
            ["Fixed 6+6", "0.394", "14.908", "0.787", "0.316"],
        ],
        ["Method", "Acc@QA", "Cost", "IoP@0.5", "Acc@GQA"],
        font_size=12,
        highlight_start=2,
    )
    add_block(slide, 0.9, 4.25, 5.45, 1.5, "Observation", "Grounding improves as evidence budget increases, but answer accuracy remains near 0.38.", accent=BLUE, body_size=18)
    add_block(slide, 6.95, 4.25, 5.45, 1.5, "Conclusion", "A weak answerer hides evidence-policy improvements in the answer metric.", accent=GOLD, body_size=18)

    # 11
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Result 2: Stronger VLM Reveals the Tradeoff", 11, total)
    add_table(
        slide,
        0.75,
        0.95,
        11.85,
        2.35,
        [
            ["Qwen one segment", "0.668", "1.500", "0.311", "0.221"],
            ["Qwen learned two-item", "0.696", "2.795", "0.410", "0.297"],
            ["Qwen fixed 3+3 reference", "0.724", "7.500", "0.615", "0.454"],
        ],
        ["Method", "Acc@QA", "Cost", "IoP@0.5", "Acc@GQA"],
        font_size=13,
        highlight_start=2,
    )
    add_metric_cards(
        slide,
        [
            ("Answer retained", "96%", GREEN),
            ("Cost used", "37%", BLUE),
            ("Grounding retained", "65%", GOLD),
        ],
        x=0.95,
        y=4.0,
        w=11.4,
        h=0.95,
    )
    add_textbox(slide, 1.1, 5.55, 11.1, 0.7, "Learned two-item evidence is the better budgeted method: near-fixed accuracy at much lower evidence cost.", font_size=20, color=BLUE_DARK, bold=True, align=PP_ALIGN.CENTER)

    # 12
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Accuracy-Cost-Grounding Tradeoff", 12, total)
    slide.shapes.add_picture(str(TRADEOFF_PATH), Inches(0.55), Inches(0.85), width=Inches(12.2))
    add_block(slide, 1.1, 5.55, 11.1, 1.12, "Interpretation", "The learned policy gives the best budgeted tradeoff: most of the high-cost answer accuracy, much lower evidence cost, and better grounding than one segment.", accent=BLUE, body_size=15)

    # 13
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Qualitative Lessons", 13, total)
    add_table(
        slide,
        0.55,
        0.9,
        12.25,
        3.55,
        [
            ["Correct and grounded", "Evidence overlaps the support moment and Qwen predicts the gold answer.", "Desired behavior"],
            ["Correct, weakly grounded", "Qwen predicts correctly but selected evidence has low temporal overlap.", "Answer accuracy alone misleads"],
            ["Grounded but wrong", "Selected evidence overlaps support span, but answer reasoning fails.", "Evidence and reasoning are separate"],
        ],
        ["Case", "What happens", "Lesson"],
        font_size=11.5,
    )
    add_block(slide, 1.0, 5.0, 11.35, 1.05, "Takeaway", "Grounded VideoQA needs both components: evidence policy must find the right moment, and the answerer must reason from it.", accent=GREEN, body_size=18)

    # 14
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Answers to the Opening Questions", 14, total)
    answers = [
        ("Q1. Same evidence budget for every question?", "No. Evidence need varies; fixed context can be wasteful or weakly grounded."),
        ("Q2. Can compact learned evidence preserve accuracy?", "Yes. Learned two-item evidence keeps about 96% of fixed 3+3 answer accuracy at 37% of the cost."),
        ("Q3. Evidence problem or answerer problem?", "Both. The frozen answerer hides evidence gains; Qwen exposes the accuracy-cost-grounding frontier."),
    ]
    for i, (q, a) in enumerate(answers):
        add_block(slide, 0.85, 0.9 + i * 1.65, 11.65, 1.18, q, a, accent=[BLUE, GREEN, GOLD][i], body_size=17)
    add_textbox(slide, 1.0, 6.15, 11.3, 0.4, "This answers the initial questions after the method and evidence-budget experiments.", font_size=16, color=GRAY, align=PP_ALIGN.CENTER)

    # 15
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Final Takeaway", 15, total)
    add_textbox(slide, 0.9, 1.0, 11.55, 0.85, "Grounded VideoQA should be evaluated as evidence selection plus answer generation.", font_size=25, color=BLUE_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_bullets(
        slide,
        1.25,
        2.35,
        10.9,
        2.2,
        [
            "A weak answerer can hide better evidence selection.",
            "A strong VLM makes compact evidence useful for answering.",
            "Faithful temporal grounding still benefits from broader evidence coverage.",
            "The result is a Pareto tradeoff, not a single winner.",
        ],
        font_size=21,
    )
    add_block(slide, 1.6, 5.35, 10.1, 0.95, "One-line closing", "Select enough evidence, not too much evidence, and verify that the answer is supported by the right moment.", accent=GREEN, body_size=18)

    prs.save(PPTX_PATH)
    prs.save(PPTX_ALIAS_PATH)


if __name__ == "__main__":
    make_figures()
    build_deck()
